"""Rebuild raid item JSON from Items only; never changes character data."""
import json,re
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
THEMES={'tactics':[5,6,7,8,9,26,35],'offense':[10,14,25,29,30,32,33],'defense':[11,15,16,21,22,24,31],'tempo':[3,4,12,13,17,18,23],'recovery':[2,19,20,27,28,34,36]}

def main():
    ws=load_workbook(ROOT/'data/Raid_Items_Source.xlsx',read_only=True,data_only=True)['Items']
    items=[]
    for row in range(2,37):
        name,effect,eligible,cost=next(ws.iter_rows(min_row=row,max_row=row,max_col=4,values_only=True))
        key=re.sub('[^a-z0-9]+','_',name.casefold()).strip('_')
        items.append(dict(key=key,name=name,effect=effect,eligible=eligible.replace('Kanbayshi','Kanbayashi'),shard_cost=int(cost),drop_rate=.005 if cost==25 else 1/(cost*5),weekly_cap=int(ws.cell(row,7).value),theme=next(t for t,rows in THEMES.items() if row in rows)))
    for item in items:
        if item['key']=='captain_s_armband':item['eligible']+=', Futakuchi'
        item['eligible']=', '.join(sorted(set(s.strip() for s in item['eligible'].split(',')),key=str.casefold))
        fixes={'death_stare':"At start of turn, inflicts Jmp-5 on the foe with the highest Jmp through its next action.",'burden_of_expectation':'Grants Pow/Qik/Spd+5. Inflicts Def/Rec-5.','water_bottle':'At start of even-numbered turns, restores 20 HP.'}
        if item['key'] in fixes:item['effect']=fixes[item['key']]
    if abs(sum(it['drop_rate'] for it in items)+.245-1)>1e-9:raise ValueError('The corrected full loot rates must sum to 100%.')
    (ROOT/'data/items.json').write_text(json.dumps(items,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    print('Wrote 35 items from Items; used corrected drop-rate formula, not original E cells.')

if __name__=='__main__':main()
