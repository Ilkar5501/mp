from __future__ import annotations

import json
from pathlib import Path
from collections import Counter
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'data' / 'HiQ Battler Abilities.xlsx'
OUT = ROOT / 'data' / 'characters.json'
CFG_OUT = ROOT / 'data' / 'game_config.json'

# These are fallback display names only. The spreadsheet always wins when a name exists.
NAME_OVERRIDES = {
    ('Washio','Silent Pillar','U'): 'Granite Drop',
    ('Komi','Energizing Spirit','U'): 'Rebound Spark',
    ('Atsumu','Arrogant Genius','U'): 'No Contest',
    ('Atsumu','Arrogant Genius','A'): 'Overclocked Rhythm',
    ('Atsumu','Arrogant Genius','B'): 'Elite Standard',
    ('Atsumu','Arrogant Genius','C'): 'Gearshift',
    ('Osamu','Twin Shadow','U'): 'Cut Through',
    ('Osamu','Twin Shadow','A'): 'Quiet Pursuit',
    ('Osamu','Twin Shadow','B'): 'Early Bite',
    ('Osamu','Twin Shadow','C'): 'Quickened Rise',
    ('Akagi','Relentless Receiver','U'): 'Full Stretch',
    ('Akagi','Relentless Receiver','A'): 'Comeback Rush',
    ('Akagi','Relentless Receiver','B'): 'Gap Sense',
    ('Akagi','Relentless Receiver','C'): 'Dead Zone',
    ('Yamamoto','Fire Lord','U'): 'Overrun',
    ('Yamamoto','Fire Lord','A'): 'Headlong Leap',
    ('Yamamoto','Fire Lord','B'): 'Pile On',
    ('Yamamoto','Fire Lord','C'): 'Passing Lift',
    ('Fukunaga','Doctor Freeze','U'): 'Encore',
    ('Fukunaga','Doctor Freeze','A'): 'Deadpan Focus',
    ('Fukunaga','Doctor Freeze','B'): 'Isolation Chamber',
    ('Fukunaga','Doctor Freeze','C'): 'Cold Silence',
    ('Yaku','Hands of Ruin','U'): 'Zero-Hour Rescue',
    ('Yaku','Hands of Ruin','A'): 'Reflected Edge',
    ('Yaku','Hands of Ruin','B'): 'Prepared for Anything',
    ('Yaku','Hands of Ruin','C'): 'Drain the Threat',
    ('Daichi','Steadfast Captain','U'): 'Keep Us Standing',
    ('Daichi','Steadfast Captain','A'): 'Refuse to Fall',
    ('Daichi','Steadfast Captain','B'): 'Built to Endure',
    ('Daichi','Steadfast Captain','C'): 'Clear the Air',
    ('Suga','Relentless Setter','U'): 'Precision Burst',
    ('Suga','Relentless Setter','A'): 'Measured Barrage',
    ('Suga','Relentless Setter','B'): 'Already There',
    ('Suga','Relentless Setter','C'): 'Pass the Torch',
    ('Tanaka','Unbreakable Wing','U'): 'All-Out Slam',
    ('Tanaka','Unbreakable Wing','A'): 'Fearless Entry',
    ('Tanaka','Unbreakable Wing','B'): 'Spreading Grit',
    ('Tanaka','Unbreakable Wing','C'): 'Keeps Burning',
    ('Oikawa','Grand Conductor','U'): 'Courtbreaker',
    ('Oikawa','Grand Conductor','A'): "Showman's Entrance",
    ('Oikawa','Grand Conductor','B'): 'Critical Cadence',
    ('Oikawa','Grand Conductor','C'): 'Skyward Command',
    ('Iwaizumi','Fitness Fanatic','U'): 'Hammer Through',
    ('Iwaizumi','Fitness Fanatic','A'): 'Meet It Head-On',
    ('Iwaizumi','Fitness Fanatic','B'): 'Immediate Reply',
    ('Iwaizumi','Fitness Fanatic','C'): 'No Half Measures',
    ('Kyotani','Mad Dog','U'): 'Unleashed Interception',
    ('Kyotani','Mad Dog','A'): "Predator's Pace",
    ('Kyotani','Mad Dog','B'): 'Feral Advantage',
    ('Kyotani','Mad Dog','C'): 'Rabid Energy',
    ('Tendo','Guess Monster','U'): 'Afterimage Curse',
    ('Tendo','Guess Monster','A'): 'Gamble on It',
    ('Tendo','Guess Monster','B'): 'Hijacked Speed',
    ('Tendo','Guess Monster','C'): 'Reverse Fortune',
    ('Yamagata','Last Line','U'): 'Nothing Gets Through',
    ('Yamagata','Last Line','A'): 'Always in Range',
    ('Yamagata','Last Line','B'): 'Brace for Impact',
    ('Yamagata','Last Line','C'): 'Emergency Formation',
    ('Reon','Benkei','B'): 'Safety Net',
    ('Reon','Benkei','C'): 'Slow the Assault',
    ('Teradomari','Agile Hunter','U'): 'Terminal Velocity',
    ('Teradomari','Agile Hunter','A'): 'Overextend',
    ('Teradomari','Agile Hunter','B'): 'Find the Opening',
    ('Teradomari','Agile Hunter','C'): 'Marked Prey',
    ('Gao','Hulking Giant','U'): 'Skyhammer',
    ('Gao','Hulking Giant','A'): 'Dual Stance',
    ('Gao','Hulking Giant','B'): 'Over the Front',
    ('Gao','Hulking Giant','C'): 'Tall Order',
    ('Kiryu','Reluctant Titan','U'): 'Pierce the Fortress',
    ('Kiryu','Reluctant Titan','A'): 'Crush the Guard',
    ('Kiryu','Reluctant Titan','B'): 'No One Steps In',
    ('Kiryu','Reluctant Titan','C'): 'Open the Wound',
    ('Unnan','Fortress of Fortitude','U'): 'Absorb the Blow',
    ('Unnan','Fortress of Fortitude','A'): 'Choke the Offense',
    ('Unnan','Fortress of Fortitude','B'): 'Dig In',
    ('Unnan','Fortress of Fortitude','C'): 'Ground Their Game',
    ('Daisho','Serpent of Chaos','U'): 'Venomous Followthrough',
    ('Daisho','Serpent of Chaos','A'): 'Sudden Lunge',
    ('Daisho','Serpent of Chaos','B'): 'Wrong-Footed',
    ('Daisho','Serpent of Chaos','C'): 'Forced Error',
}


def clean_num(v):
    if v is None:
        return 0
    try:
        iv = int(v)
        return iv
    except (TypeError, ValueError):
        return v


def main():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    ws = wb['Characters']
    headers = [c.value for c in ws[1]]
    h = {name: i for i, name in enumerate(headers) if name}

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[h['Name:']]:
            continue
        name = str(r[h['Name:']]).strip()
        title = str(r[h['Title:']] or '').strip()
        char_id = str(r[h['Character ID']] or (f'{name} — {title}' if title else name)).strip()
        slots = {}
        for slot, name_col, text_col, timing_col in [
            ('U','Ultimate Name:','Ultimate:','Ultimate Timing'),
            ('A','A Name:','A Skill:','A Timing'),
            ('B','B Name:','B Skill:','B Timing'),
            ('C','C Name:','C Skill:','C Timing'),
        ]:
            ability_name = str(r[h[name_col]] or '').strip()
            if not ability_name:
                ability_name = NAME_OVERRIDES.get((name,title,slot), '')
            slots[slot] = {
                'name': ability_name,
                'text': str(r[h[text_col]] or '').strip(),
                'timing': str(r[h[timing_col]] or '').strip(),
            }
        rows.append({
            'id': char_id,
            'name': name,
            'title': title,
            'position': str(r[h['Position:']] or '').strip(),
            'school': str(r[h['School:']] or '').strip(),
            'rarity': clean_num(r[h['Rarity:']]),
            'role_class': str(r[h.get('Class:', -1)] or '').strip() if 'Class:' in h else '',
            'card_art': str(r[h['Card Art:']] or '').strip(),
            'character_icon': str(r[h['Character Icon:']] or '').strip(),
            'stats': {
                'hp': clean_num(r[h['HP:']]),
                'pow': clean_num(r[h['POW:']]),
                'qik': clean_num(r[h['QIK:']]),
                'def': clean_num(r[h['DEF:']]),
                'rec': clean_num(r[h['REC:']]),
                'spd': clean_num(r[h['SPD:']]),
                'jmp': clean_num(r[h['JMP:']]),
                'bst': clean_num(r[h['BST:']]),
            },
            'ult_cd': clean_num(r[h['Ult CD:']]),
            'abilities': slots,
            'extended': {
                'U': str(r[h['Extended Skill Info U:']] or '').strip() if 'Extended Skill Info U:' in h else '',
                'A': str(r[h['Extended Skill Info A:']] or '').strip(),
                'B': str(r[h['Extended Skill Info B:']] or '').strip(),
                'C': str(r[h['Extended Skill Info C:']] or '').strip(),
            },
            'clarification_count': clean_num(r[h['Clarification Count']]),
        })

    # Keep only named status definitions as extended tooltips.
    for character in rows:
        for slot,text in character['extended'].items():
            if text and not text.startswith('['):character['extended'][slot]=''
        if character['name']=='Kiryu' and character['title']=='Last Hope':
            character['abilities']['C']['text']=character['abilities']['C']['text'].replace('[Burden]','Burden')
    from apply_character_updates import apply_updates
    rows = apply_updates(rows)
    OUT.write_text(json.dumps(rows, indent=2, ensure_ascii=False), encoding='utf-8')

    # Pack rules are intentionally explicit rather than parsed from prose so the bot is deterministic.
    school_counts = Counter(x['school'] for x in rows)
    eligible_schools = sorted([school for school, count in school_counts.items() if count >= 5])
    cfg = {
        'prefix': ';',
        'timezone': 'America/New_York',
        'packs': {
            'base': {
                'cost': 300,
                'cards': 3,
                'rarity_odds': {'1': 30, '2': 25, '3': 20, '4': 15, '5': 10},
            },
            'position': {
                'cost': 450,
                'cards': 3,
                'rarity_odds': {'1': 20, '2': 20, '3': 25, '4': 20, '5': 15},
            },
            'school': {
                'cost': 450,
                'cards': 3,
                'rarity_odds': {'1': 20, '2': 20, '3': 25, '4': 20, '5': 15},
                'eligible_schools': eligible_schools,
            },
        },
        'daily': {
            'free_base_packs': 1,
            'win_rewards': {'1': 50, '2': 100, '3': 150},
        },
        'weekly': {
            'win_milestones': {'10': 300, '20': 300, '30': 300, '40': 300, '50': 300},
        },
        'gauntlet': {'clear_reward': 450},
        'merge': {
            'max_level': 10,
            'stat_cycle': ['offense', 'def', 'rec', 'spd', 'jmp'],
            'hp_milestones': [5, 10],
        },
        'boon_bane': {
            'magnitude': 3,
            'eligible_stats': ['hp', 'offense', 'def', 'rec', 'spd', 'jmp'],
            'neutral_when_same': True,
        },
        'battle': {
            'max_turns': 30,
            'base_target_weights': {'front': 40, 'back': 30},
            'crit_per_jmp_difference_percent': 5,
            'crit_multiplier': 1.2,
            'school_synergy_per_ally': 1,
            'school_synergy_cap': None,
        },
    }
    CFG_OUT.write_text(json.dumps(cfg, indent=2, ensure_ascii=False), encoding='utf-8')
    print(f'Wrote {len(rows)} characters to {OUT}')
    print('School packs:', ', '.join(eligible_schools))

if __name__ == '__main__':
    main()
